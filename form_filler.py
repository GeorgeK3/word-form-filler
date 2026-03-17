from docx import Document
import tkinter as tk
from tkinter import messagebox
import re

# --- Edit function ---
def edit_cell_below(doc, search_value, new_value):
    for i, table in enumerate(doc.tables):
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                if cell.text.strip() == search_value:
                    if r + 1 < len(table.rows):
                        table.cell(r + 1, c).text = new_value
                        print(f"OK Table {i} ({r},{c}): set below '{search_value}' to '{new_value}'")

# --- Validation functions ---
def is_valid_name(value):
    return value.strip().replace(" ", "").isalpha()
    

def is_valid_age(value):
    return value.strip().isdigit() and 1 <= int(value.strip()) <= 120

def is_valid_car_id(value):
    pattern = r'^[A-ZΑ-Ω]{3}-\d{4}$'
    return bool(re.match(pattern, value.strip().upper()))

# --- Highlight field red/green ---
def mark_field(entry, is_valid):
    entry.config(bg="lightgreen" if is_valid else "lightcoral")



# --- Submit action ---
def submit():
    doc = Document("file.docx")

    name   = name_entry.get()
    last_name = last_name_entry.get()
    age    = age_entry.get()
    car_id = car_entry.get()

    # Validate each field
    name_ok   = is_valid_name(name)
    last_name_ok = is_valid_name(last_name)
    age_ok    = is_valid_age(age)
    car_ok    = is_valid_car_id(car_id)

    # Color the fields
    mark_field(name_entry,  name_ok)
    mark_field(last_name_entry,  last_name_ok)
    mark_field(age_entry,   age_ok)
    mark_field(car_entry,   car_ok)

    # Stop if anything is invalid
    if not all([name_ok, last_name_ok, age_ok, car_ok]):
        errors = []
        if not name_ok:         errors.append("• Name: letters only")
        if not last_name_ok:    errors.append("• Last Name: letters only")
        if not age_ok:          errors.append("• Age: number between 1-120")
        if not car_ok:          errors.append("• Car ID: format must be AAA-1234")
        messagebox.showerror("Invalid Input", "\n".join(errors))
        return

    # All good — edit and save
    edit_cell_below(doc, "Name",      name)
    edit_cell_below(doc, "Last Name", last_name)
    edit_cell_below(doc, "Age",       age)
    edit_cell_below(doc, "Car_Id",    car_id)

    doc.save(last_name+".docx")
    messagebox.showinfo("Success", "File saved as " + last_name + ".docx!")


# --- Build the GUI ---
root = tk.Tk()
root.title("Form Filler")
root.geometry("300x350")

# Name field
tk.Label(root, text="Name").pack(pady=(30, 0))
name_entry = tk.Entry(root, width=30)
name_entry.pack()

# Last Name field
tk.Label(root, text="Last Name").pack(pady=(10, 0))
last_name_entry = tk.Entry(root, width=30)
last_name_entry.pack()

# Age field
tk.Label(root, text="Age").pack(pady=(10, 0))
age_entry = tk.Entry(root, width=30)
age_entry.pack()

# Car ID ty field
tk.Label(root, text="Car_Id").pack(pady=(10, 0))
car_entry = tk.Entry(root, width=30)
car_entry.pack()

# Submit button
tk.Button(root, text="Fill & Save", command=submit).pack(pady=20)

# Bind Enter key to submit
root.bind("<Return>", lambda event: submit())

# Define field order
fields = [name_entry, last_name_entry, age_entry, car_entry]

def focus_prev(event):
    current = root.focus_get()
    idx = fields.index(current)
    fields[(idx - 1) % len(fields)].focus_set()  # wrap around

def focus_next(event):
    current = root.focus_get()
    idx = fields.index(current)
    fields[(idx + 1) % len(fields)].focus_set()  # wrap around

root.bind("<Up>",   focus_prev)
root.bind("<Down>", focus_next)

root.mainloop()



